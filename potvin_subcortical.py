"""Four Potvin mmc2 models, numerically cross-checked with LibreOffice Calc.

Research normative comparison only, not native Excel or clinical validation.
"""
import hashlib
import math
import warnings
from pathlib import Path
import numpy as np
import pandas as pd
from scipy.stats import t

MODELS = {'Left_Hippocampus': 18, 'Right_Hippocampus': 19,
          'Left_Amygdala': 13, 'Right_Amygdala': 14}
WORKBOOK_SHA256 = 'c8a8c9e08d5703923a45eefe2d29ef6dc9c553fc12771f2f35df0af872bfc19f'
SOURCE_DOI = '10.1016/j.neuroimage.2016.05.016'
SOURCE_TITLE = 'Potvin et al. (2016): Normative data for subcortical regional volumes over the lifetime of the adult human brain'


def extract(workbook):
    import openpyxl
    # Pin the complete formula/parameter source, not only selected formula substrings.
    if hashlib.sha256(Path(workbook).read_bytes()).hexdigest() != WORKBOOK_SHA256:
        raise ValueError('Unverified mmc2 workbook revision; validate before use')
    # We only read formulas and never save this workbook; unsupported display
    # extensions cannot be lost on disk. Do not show a misleading save warning.
    with warnings.catch_warnings():
        warnings.filterwarnings('ignore', message='Unknown extension is not supported and will be removed',
                                category=UserWarning, module='openpyxl.worksheet._reader')
        w = openpyxl.load_workbook(workbook, keep_vba=True)
    if w['Statistics']['A1'].value != 'SUBCORTICAL NORMS CALCULATOR':
        raise ValueError('Unexpected calculator')
    if w['Matrix']['F2'].value != '=Statistics!F6- 47.5634617' or w['Matrix']['J2'].value != '=Statistics!J6-1521907.28':
        raise ValueError('Unexpected centering')
    def named(name):
        sheet, area = next(w.defined_names[name].destinations)
        return [[c.value for c in row] for row in w[sheet][area]]
    result = []
    for name, row in MODELS.items():
        loc = next(i for i in range(3, 30) if w['Matrix'].cell(i, 1).value == name)
        formula = w['Matrix'].cell(loc, 2).value
        if formula != f'=TINV(0.05, C{loc}-11)':
            raise ValueError('Unexpected critical t formula')
        formulas = {c: getattr(w['Statistics'][f'{c}{row}'].value, 'text', w['Statistics'][f'{c}{row}'].value) for c in ('C','D','E')}
        if f'MMULT(X_{name}, B_{name})' not in formulas['C']:
            raise ValueError('Unsupported prediction formula')
        terms = named('pred_' + name)[0]
        beta = [r[0] for r in named('B_' + name)]
        covariance = named('M_' + name)
        if len(terms) != len(beta) or np.asarray(covariance).shape != (len(beta), len(beta)):
            raise ValueError('Model dimensions mismatch')
        result.append(dict(name=name, terms=terms, beta=beta, covariance=covariance,
            n=w['Matrix'].cell(loc, 3).value, mse=w['Matrix'].cell(loc, 4).value,
            critical_df=w['Matrix'].cell(loc, 3).value-11, formulas=formulas,
            source_cells=f'Statistics!C{row}:E{row}'))
    w.close()
    return result


def predict(model, patient):
    age = float(patient['age']); etiv = float(patient['estimated_total_intracranial_volume'])
    field = float(patient['scanner_field_strength'])
    sex = {'male': 1., 'female': 0.}[str(patient['sex']).lower()]
    manufacturer = str(patient['scanner_manufacturer']).lower()
    if manufacturer not in ('siemens', 'ge', 'philips') or field not in (1.5, 3.):
        raise ValueError('Unsupported scanner encoding')
    if not 18 <= age <= 94 or not math.isfinite(etiv) or etiv <= 0:
        raise ValueError('Age outside model coverage or invalid eTIV')
    a, v, f = age-47.5634617, etiv-1521907.28, float(field == 1.5)
    ge, philips = float(manufacturer == 'ge'), float(manufacturer == 'philips')
    values = dict(b0=1., agec=a, agec2=a*a, agec3=a*a*a, sexq=sex, tivc=v,
        tivc2=v*v, tivc3=v*v*v, mfsq=f, ge=ge, philips=philips,
        ge_x_mfsq=ge*f, philips_x_mfsq=philips*f, tivc_x_mfsq=v*f,
        agec_x_sexq=a*sex, tivc_x_ge=v*ge, tivc_x_philips=v*philips)
    x = np.array([values[name.lower()] for name in model['terms']])
    mean = float(x @ np.asarray(model['beta']))
    variance = model['mse'] * (1 + x @ np.asarray(model['covariance']) @ x)
    if not math.isfinite(mean) or mean <= 0 or not math.isfinite(variance) or variance <= 0:
        raise ValueError('Invalid predictive variance')
    width = t.ppf(.975, model['critical_df']) * math.sqrt(variance)
    return mean, mean-width, mean+width


def preview(result, sid):
    from comparison_atlas import comparison_features, translate_records, feature_key, mapping_error
    from mmc_normative_workflow import classify_prediction_interval
    from verify_mmc2_fixture import verify
    workbook = Path(result['project_dir']) / 'workflow_sources/external_validators/potvin_subcortical/mmc2.xlsm'
    models = extract(workbook)
    validation = verify(result['project_dir'], models)
    frame = comparison_features(result)
    patient = frame.loc[frame.subject_id.eq(sid)]
    records = []
    sha = hashlib.sha256(workbook.read_bytes()).hexdigest()
    for m in models:
        roi = m['name'].replace('_', '-')
        reference = translate_records([dict(roi_name=roi, atlas_name='FreeSurfer aseg',
            hemisphere='lh' if roi.startswith('Left-') else 'rh', imaging_metric='roi_volume')],
            frame.iloc[0]['comparison_registry'] if len(frame) else None)[0]
        matches = patient.loc[[feature_key(p) == feature_key(reference) and not mapping_error(p)
                               for p in patient.to_dict('records')]]
        row = dict(subject_id=sid, roi_name=roi, hemisphere=reference.get('hemisphere'),
            imaging_metric='roi_volume', unit='mm3', observed=None, expected=None,
            lower_95pi=None, upper_95pi=None, source_cells=m['source_cells'],
            source_doi=SOURCE_DOI, source_title=SOURCE_TITLE, workbook_sha256=sha,
            status='unavailable', range_status='not_assessed',
            validation_status=validation['status'],
            limitation='Research comparison: Calc formula concordance passed, not native Excel or clinical validation. Segmentation QC remains required; pointwise PI, no multiplicity correction.')
        try:
            if mapping_error(reference):
                raise ValueError(mapping_error(reference))
            if len(matches) != 1 or matches.iloc[0]['unit'] != 'mm3':
                raise ValueError('Missing or ambiguous matching mm3 patient volume')
            p = matches.iloc[0].to_dict()
            row.update(roi_name=p['roi_name'], hemisphere=p['hemisphere'])
            if Path(str(p.get('source_file', ''))).name != 'aseg.stats':
                raise ValueError('Native aseg.stats source required for mmc2 comparison')
            versions = [str(p.get(k, '')).strip() for k in ('processing_software_version', 'atlas_version')]
            versions = [v for v in versions if v and v.lower() not in ('nan', 'none')]
            if not versions or any(v not in ('5.3', '5.3.0') for v in versions):
                raise ValueError('Subcortical preview requires consistent explicit FreeSurfer 5.3 provenance')
            observed = float(p['value_numeric'])
            if not math.isfinite(observed) or observed < 0:
                raise ValueError('Invalid measured volume')
            mean, lower, upper = predict(m, p)
            prediction_se = (upper-mean)/t.ppf(.975, m['critical_df'])
            row.update(observed=observed, expected=mean, lower_95pi=lower, upper_95pi=upper,
                       patient_processing_version=versions[0], status='calculated',
                       range_status=classify_prediction_interval(observed, lower, upper),
                       zop=(observed-mean)/math.sqrt(m['mse']),
                       prediction_se=prediction_se,
                       percentile=100*t.cdf((observed-mean)/prediction_se, m['critical_df']),
                       model_id='potvin2016_mmc2_' + m['name'])
        except (ValueError, TypeError, KeyError) as exc:
            row.update(status='unavailable', limitation=str(exc))
        records.append(row)
    return pd.DataFrame(records)


def augment_normative(result, normative):
    """Replace only the four matched native measurements, never add/double count rows."""
    data = normative.copy()
    missing = data.exclusion_reason.eq('No matching cortical DK model in mmc1 (aseg/subcortical not covered)')
    data.loc[missing, 'exclusion_reason'] = ('No implemented compatible Potvin model for this feature '
        '(mmc1 cortical; mmc2 currently covers bilateral hippocampus/amygdala only)')
    for sid in data.subject_id.unique():
        try:
            rows = preview(result, sid)
        except (OSError, ValueError, KeyError) as exc:
            # Keep cortical output available, but make a broken validator visible.
            mask = data.roi_name.isin([m.replace('_', '-') for m in MODELS]) & data.subject_id.eq(sid)
            data.loc[mask, 'exclusion_reason'] = 'mmc2 validation unavailable: ' + str(exc)
            continue
        for row in rows.to_dict('records'):
            mask = (data.subject_id.eq(sid) & data.roi_name.eq(row['roi_name']) &
                    data.hemisphere.eq(row['hemisphere']) & data.imaging_metric.eq('roi_volume'))
            if int(mask.sum()) != 1:
                continue
            values = dict(source_title=SOURCE_TITLE, source_doi=SOURCE_DOI,
                          source_sha256=row['workbook_sha256'], source_cells=row['source_cells'],
                          method_note=row['limitation'], validation_status=row['validation_status'],
                          normative_review_status='pending_human_review')
            if row['status'] == 'calculated':
                values.update(calculation_status='calculated', expected_value=row['expected'],
                    lower_95pi=row['lower_95pi'], upper_95pi=row['upper_95pi'],
                    range_status=row['range_status'], zop=row['zop'], percentile=row['percentile'],
                    prediction_se=row['prediction_se'],
                    expected_direction=('below_expected' if row['observed'] < row['expected'] else
                                        'above_expected' if row['observed'] > row['expected'] else 'at_expected'),
                    model_id=row['model_id'], exclusion_reason='')
            else:
                values.update(calculation_status='not_calculated', exclusion_reason=row['limitation'])
            for key, value in values.items():
                data.loc[mask, key] = value
    return data
